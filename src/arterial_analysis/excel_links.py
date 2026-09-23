"""저장된 Excel 셀값과 외부참조식을 다루는 연결 계층."""

from __future__ import annotations

import re
from pathlib import Path


SUPPORTED_EXTENSIONS = {".xlsx", ".xls", ".xlsm", ".xlsb"}
OPENPYXL_EXTENSIONS = {".xlsx", ".xlsm"}
_EXTERNAL_FORMULA = re.compile(
    r"^=\s*'?(?P<prefix>.*?\[)(?P<book>[^\]]+)\](?P<sheet>[^']+)'?!\$?(?P<column>[A-Za-z]{1,3})\$?(?P<row>[1-9]\d*)$"
)
_LOCAL_FORMULA = re.compile(r"^=\s*'?(?P<sheet>[^']+)'?!\$?(?P<column>[A-Za-z]{1,3})\$?(?P<row>[1-9]\d*)$")


def _excel_modules():
    import pythoncom
    import win32com.client
    return pythoncom, win32com.client


def normalize_cell(address: str) -> str:
    return str(address or "").replace("$", "").strip().upper()


def format_external_formula(path: str, sheet: str, address: str) -> str:
    """Excel과 같은 외부 통합문서 참조식으로 표시한다."""
    source = Path(path).resolve()
    cell = normalize_cell(address)
    match = re.fullmatch(r"([A-Z]{1,3})([1-9]\d*)", cell)
    if not match:
        raise ValueError("Excel 셀 주소가 올바르지 않습니다.")
    escaped_sheet = str(sheet).replace("'", "''")
    return f"='{source.parent}\\[{source.name}]{escaped_sheet}'!${match.group(1)}${match.group(2)}"


def parse_external_formula(formula: str, default_path: str = "") -> tuple[str, str, str]:
    """전체 외부참조식 또는 같은 파일의 시트 참조식을 해석한다."""
    text = str(formula or "").strip()
    match = _EXTERNAL_FORMULA.fullmatch(text)
    if match:
        prefix = match.group("prefix")[:-1].replace("''", "'")
        path = str((Path(prefix) / match.group("book")).resolve())
        sheet = match.group("sheet").replace("''", "'")
        return path, sheet, normalize_cell(match.group("column") + match.group("row"))
    match = _LOCAL_FORMULA.fullmatch(text)
    if match and default_path:
        return str(Path(default_path).resolve()), match.group("sheet").replace("''", "'"), normalize_cell(match.group("column") + match.group("row"))
    raise ValueError("연결식은 ='폴더\\[파일.xlsx]시트'!$A$1 형식이어야 합니다.")


def _openpyxl_book(path: str):
    from openpyxl import load_workbook
    source = Path(path).resolve()
    return load_workbook(source, read_only=True, data_only=True, keep_links=False, keep_vba=source.suffix.lower() == ".xlsm")


def list_sheets(path: str) -> list[str]:
    if Path(path).suffix.lower() in OPENPYXL_EXTENSIONS:
        workbook = _openpyxl_book(path)
        try:
            return list(workbook.sheetnames)
        finally:
            workbook.close()
    pythoncom, client = _excel_modules(); pythoncom.CoInitialize()
    excel = None; workbook = None
    try:
        excel = client.DispatchEx("Excel.Application")
        excel.Visible = False; excel.DisplayAlerts = False; excel.AskToUpdateLinks = False
        workbook = excel.Workbooks.Open(str(Path(path).resolve()), UpdateLinks=0, ReadOnly=True)
        return [str(sheet.Name) for sheet in workbook.Worksheets]
    finally:
        if workbook is not None: workbook.Close(SaveChanges=False)
        if excel is not None: excel.Quit()
        pythoncom.CoUninitialize()


def read_saved_cell(path: str, sheet: str, address: str):
    """디스크에 저장된 통합문서에서 단일 셀값을 읽는다."""
    result = read_saved_cells(path, sheet, [address])[normalize_cell(address)]
    if result[2]: raise ValueError(result[2])
    return result[0], result[1]


def read_saved_cells(path: str, sheet: str, addresses: list[str]):
    """통합문서를 한 번만 열어 같은 시트의 여러 셀을 일괄 읽는다."""
    if Path(path).suffix.lower() in OPENPYXL_EXTENSIONS:
        workbook = _openpyxl_book(path)
        try:
            if str(sheet) not in workbook.sheetnames:
                raise ValueError(f"'{sheet}' 시트를 찾을 수 없습니다.")
            worksheet = workbook[str(sheet)]; results = {}
            for address in dict.fromkeys(normalize_cell(a) for a in addresses):
                try:
                    cell = worksheet[address]
                    merged_cells = getattr(worksheet, "merged_cells", ())
                    if cell.coordinate in merged_cells:
                        merged = next(item for item in merged_cells.ranges if cell.coordinate in item)
                        cell = worksheet.cell(merged.min_row, merged.min_col)
                    value = cell.value
                    if value is None or isinstance(value, (str, bool)):
                        raise ValueError("원본 셀이 비어 있거나 저장된 숫자값이 아닙니다.")
                    results[address] = (int(round(float(value))), normalize_cell(cell.coordinate), "")
                except Exception as exc:
                    results[address] = (None, address, str(exc))
            return results
        finally:
            workbook.close()
    pythoncom, client = _excel_modules(); pythoncom.CoInitialize()
    excel = None; workbook = None
    try:
        excel = client.DispatchEx("Excel.Application")
        excel.Visible = False; excel.DisplayAlerts = False; excel.AskToUpdateLinks = False
        workbook = excel.Workbooks.Open(str(Path(path).resolve()), UpdateLinks=0, ReadOnly=True)
        worksheet=workbook.Worksheets(str(sheet));results={}
        for address in dict.fromkeys(normalize_cell(a) for a in addresses):
            try:
                target=worksheet.Range(address)
                if bool(target.MergeCells):target=target.MergeArea.Cells(1,1)
                value=target.Value2
                if value is None or isinstance(value,(str,bool)):raise ValueError("원본 셀이 비어 있거나 숫자가 아닙니다.")
                results[address]=(int(round(float(value))),normalize_cell(str(target.Address)),"")
            except Exception as exc:results[address]=(None,address,str(exc))
        return results
    finally:
        if workbook is not None: workbook.Close(SaveChanges=False)
        if excel is not None: excel.Quit()
        pythoncom.CoUninitialize()


def open_and_activate(path: str, sheet: str = "") -> None:
    """기존 Excel을 재사용해 원본 파일을 열고 앞으로 가져온다."""
    pythoncom, client = _excel_modules(); pythoncom.CoInitialize()
    excel = client.Dispatch("Excel.Application")
    full = str(Path(path).resolve()).lower()
    workbook = next((book for book in excel.Workbooks if str(book.FullName).lower() == full), None)
    if workbook is None: workbook = excel.Workbooks.Open(str(Path(path).resolve()), UpdateLinks=0, ReadOnly=False)
    excel.Visible = True; workbook.Activate()
    if sheet: workbook.Worksheets(str(sheet)).Activate()
    excel.WindowState = -4143
    pythoncom.CoUninitialize()


def active_selection() -> tuple[str, str, str, bool]:
    """활성 Excel 선택의 파일, 시트, 셀 주소, 저장 여부를 반환한다."""
    pythoncom, client = _excel_modules(); pythoncom.CoInitialize()
    try:
        excel = client.GetActiveObject("Excel.Application")
        workbook = excel.ActiveWorkbook; sheet = excel.ActiveSheet; target = excel.Selection
        if workbook is None or sheet is None or target is None: raise RuntimeError("Excel에서 셀을 선택해 주세요.")
        if bool(target.MergeCells): target = target.MergeArea.Cells(1, 1)
        return str(workbook.FullName), str(sheet.Name), normalize_cell(str(target.Address)), bool(workbook.Saved)
    finally:
        pythoncom.CoUninitialize()
