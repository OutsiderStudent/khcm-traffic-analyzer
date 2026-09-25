"""교차로 분석 정리 Excel에서 상황별 서비스수준 참고값을 읽는다."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path


VALID_LOS = {"A", "B", "C", "D", "E", "F", "FF", "FFF"}


@dataclass(frozen=True, slots=True)
class IntersectionLosRecord:
    point_number: str
    name: str
    los: str
    cell: str


@dataclass(frozen=True, slots=True)
class IntersectionLosGroup:
    scenario: str
    year: int
    header: str
    sheet: str
    records: tuple[IntersectionLosRecord, ...]


def normalize_point(value: object) -> str:
    text = str(value or "").strip().upper()
    if not text:
        return ""
    match = re.search(r"([A-Z]|\d+)", text)
    if not match:
        return text
    token = match.group(1)
    return str(int(token)) if token.isdigit() else token


def normalize_name(value: object) -> str:
    text = re.sub(r"\s+", "", str(value or "").strip())
    return re.sub(r"^(?:\d+|[A-Za-z])\s*[.．·:-]\s*", "", text)


def parse_scenario_header(value: object) -> tuple[str, int] | None:
    text = re.sub(r"\s+", "", str(value or ""))
    if "비교" in text:
        return None
    year_match = re.search(r"(20\d{2})", text)
    if not year_match:
        return None
    if "미시행" in text:
        scenario = "사업 미시행시"
    elif "개선" in text:
        scenario = "개선대책 이행시"
    elif "시행" in text:
        scenario = "사업 시행시"
    elif "현황" in text:
        scenario = "현황"
    else:
        return None
    return scenario, int(year_match.group(1))


def read_intersection_los(path: str | Path) -> list[IntersectionLosGroup]:
    from openpyxl import load_workbook

    source = Path(path).resolve()
    workbook = load_workbook(source, data_only=True, read_only=False, keep_links=False)
    try:
        worksheet = workbook["교차로"] if "교차로" in workbook.sheetnames else None
        if worksheet is None:
            for candidate in workbook.worksheets:
                if any(str(cell.value or "").strip() == "교차로명" for row in candidate.iter_rows(min_row=1, max_row=min(12, candidate.max_row)) for cell in row):
                    worksheet = candidate
                    break
        if worksheet is None:
            raise ValueError("'교차로' 시트 또는 '교차로명' 표 머리글을 찾을 수 없습니다.")

        header_cell = next(
            (cell for row in worksheet.iter_rows(min_row=1, max_row=min(12, worksheet.max_row)) for cell in row if str(cell.value or "").strip() == "교차로명"),
            None,
        )
        if header_cell is None:
            raise ValueError("'교차로명' 열을 찾을 수 없습니다.")
        header_row = header_cell.row
        name_column = header_cell.column
        point_column = next(
            (cell.column for cell in worksheet[header_row] if str(cell.value or "").strip() in {"지점", "지점번호"}),
            None,
        )
        if point_column is None:
            raise ValueError("'지점' 열을 찾을 수 없습니다.")

        groups: list[IntersectionLosGroup] = []
        seen: set[tuple[str, int]] = set()
        for column in range(1, worksheet.max_column + 1):
            if str(worksheet.cell(header_row + 1, column).value or "").strip().upper() != "LOS":
                continue
            group_header = next(
                (worksheet.cell(header_row, left).value for left in range(column, max(0, column - 3), -1) if worksheet.cell(header_row, left).value not in (None, "")),
                None,
            )
            parsed = parse_scenario_header(group_header)
            if parsed is None or parsed in seen:
                continue
            records: list[IntersectionLosRecord] = []
            for row in range(header_row + 3, worksheet.max_row + 1):
                name = str(worksheet.cell(row, name_column).value or "").strip()
                point = normalize_point(worksheet.cell(row, point_column).value)
                los = str(worksheet.cell(row, column).value or "").strip().upper()
                if not name and not point:
                    continue
                if los not in VALID_LOS:
                    continue
                records.append(IntersectionLosRecord(point, name, los, worksheet.cell(row, column).coordinate))
            if records:
                seen.add(parsed)
                groups.append(IntersectionLosGroup(parsed[0], parsed[1], str(group_header), worksheet.title, tuple(records)))
        if not groups:
            raise ValueError("상황·연도별 LOS 열을 찾을 수 없습니다.")
        return groups
    finally:
        workbook.close()


def match_intersection_los(records: tuple[IntersectionLosRecord, ...], number: object, name: object) -> IntersectionLosRecord | None:
    normalized_name = normalize_name(name)
    if normalized_name:
        matches = [record for record in records if normalize_name(record.name) == normalized_name]
        if len(matches) == 1:
            return matches[0]
    normalized_point = normalize_point(number)
    if normalized_point:
        matches = [record for record in records if record.point_number == normalized_point]
        if len(matches) == 1:
            return matches[0]
    return None
