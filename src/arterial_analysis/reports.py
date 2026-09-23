from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
import re
from typing import Iterable

from PySide6.QtCore import QRect, Qt
from PySide6.QtGui import QColor, QFont, QImage, QPainter, QPen

from .engine import AnalysisResult
from .project import ArterialProject


@dataclass(slots=True)
class ReportSpec:
    title: str
    groups: list[tuple[str, list[str]]]
    rows: list[list[str]]
    widths: list[int]
    row_uids: list[str | None] = field(default_factory=list)
    row_categories: list[str] = field(default_factory=list)
    highlight_cells: set[tuple[int, int]] = field(default_factory=set)
    row_details: list[str] = field(default_factory=list)
    row_warnings: list[str] = field(default_factory=list)
    header_spans: list[tuple[int, int, str]] = field(default_factory=list)
    row_pair_ids: list[str] = field(default_factory=list)
    body_merge_columns: tuple[int, ...] = ()

    @property
    def headers(self) -> list[str]:
        return [header for _, headers in self.groups for header in headers]


def _key(result: AnalysisResult) -> tuple[str, str]:
    return result.segment.comparison_id, result.segment.direction


def _key_order(key: tuple[str, str]) -> tuple[str, int]:
    return key[0], 0 if key[1] == "→" else 1


def _road_group_order(category: str) -> int:
    """외부도로 계열을 먼저, 사업지 내부도로를 나중에 표시한다."""
    return 1 if category == "사업지 내부도로" else 0


def result_sort_key(result: AnalysisResult) -> tuple[int, str, str, int]:
    segment = result.segment
    return (
        _road_group_order(segment.road_category),
        segment.road_name,
        segment.comparison_id,
        0 if segment.direction == "→" else 1,
    )


def _delta_text(value: float, decimals: int = 0) -> str:
    """변화표는 양수의 +를 생략하고 정확한 0은 대시로 표시한다."""
    if abs(value) < 10 ** (-(decimals + 2)):
        return "-"
    return _fmt_number(value, decimals)


def _segment_label(result: AnalysisResult) -> str:
    s = result.segment
    start = f"{s.start_number}{s.start_name}" if str(s.start_number).strip() else s.start_name
    end = f"{s.end_number}{s.end_name}" if str(s.end_number).strip() else s.end_name
    if s.direction == "←":
        return f"{end}  ←  {start}"
    return f"{start}  →  {end}"


def _segment_cells(result: AnalysisResult) -> list[str]:
    """보고서 구간도: 양방향 공통 교차로 셀과 방향 셀을 분리한다."""
    s=result.segment
    return [s.road_name,str(s.start_number),s.start_name,s.direction,str(s.end_number),s.end_name]


SEGMENT_GROUP=("구간명",["구 분","구 간","","","",""])
SEGMENT_HEADER_SPANS=[(1,5,"구 간")]
SEGMENT_MERGE_COLUMNS=(0,1,2,4,5)
SEGMENT_WIDTHS=[100,42,130,44,42,130]
RESULT_WIDTH=100


def _report_length(result: AnalysisResult) -> float:
    value = result.segment.report_length_km
    return result.segment.length_km if value is None else value


def _report_volume(result: AnalysisResult) -> float:
    value = result.segment.report_volume
    return result.segment.main_volume if value is None else value


def _fmt_number(value: float, decimals: int = 0) -> str:
    return f"{value:,.{decimals}f}"


def current_report(project: ArterialProject, results: Iterable[AnalysisResult]) -> ReportSpec:
    current = sorted(
        (r for r in results if r.segment.scenario == "현황" and r.segment.year == project.current_year),
        key=result_sort_key,
    )
    rows = [
        [
            *_segment_cells(r),
            r.arterial_type.replace("유형 ", "유형"),
            _fmt_number(_report_length(r), 1),
            _fmt_number(_report_volume(r)),
            _fmt_number(r.speed_kmh, 1),
            r.los,
        ]
        for r in current
    ]
    return ReportSpec(
        title=f"{project.current_year}년 도시·교외간선도로 서비스수준 분석결과",
        groups=[
            SEGMENT_GROUP,
            ("분석결과", ["도로유형", "구간거리(km)", "교통량(대/시)", "평균통행속도(km/h)", "서비스수준(LOS)"]),
        ],
        rows=rows,
        widths=[*SEGMENT_WIDTHS,*([RESULT_WIDTH]*5)],
        row_uids=[r.segment.uid for r in current],
        row_categories=[r.segment.road_category for r in current],
        highlight_cells={(index, 9) for index, r in enumerate(current) if r.segment.manual_speed_kmh is not None},
        header_spans=SEGMENT_HEADER_SPANS,row_pair_ids=[r.segment.comparison_id for r in current],body_merge_columns=SEGMENT_MERGE_COLUMNS,
    )


def future_report(project: ArterialProject, results: Iterable[AnalysisResult], scenario: str) -> ReportSpec:
    future = [r for r in results if r.segment.scenario == scenario and r.segment.year in project.future_years]
    years = sorted({r.segment.year for r in future})
    lookup = {(r.segment.year, *_key(r)): r for r in future}
    first_for_key = {
        key: sorted((r for r in future if _key(r) == key), key=lambda r: r.segment.year)[0]
        for key in {_key(r) for r in future}
    }
    base_keys = sorted(first_for_key, key=lambda key: result_sort_key(first_for_key[key]))
    rows: list[list[str]] = []
    for comparison_id, direction in base_keys:
        candidates = [r for r in future if _key(r) == (comparison_id, direction)]
        if not candidates:
            continue
        first = sorted(candidates, key=lambda r: r.segment.year)[0]
        row = _segment_cells(first)
        for year in years:
            result = lookup.get((year, comparison_id, direction))
            row.extend(
                ["-", "-", "-"]
                if result is None
                else [_fmt_number(_report_volume(result)), _fmt_number(result.speed_kmh, 1), result.los]
            )
        rows.append(row)
    groups = [SEGMENT_GROUP] + [
        (f"{year}년", ["교통량\n(대/시)", "평균통행속도\n(km/h)", "서비스수준\n(LOS)"])
        for year in years
    ]
    widths = SEGMENT_WIDTHS + [RESULT_WIDTH] * (3*len(years))
    highlights: set[tuple[int, int]] = set()
    for row_index, (comparison_id, direction) in enumerate(base_keys):
        for year_index, year in enumerate(years):
            result = lookup.get((year, comparison_id, direction))
            if result and result.segment.manual_speed_kmh is not None:
                highlights.add((row_index, 7 + year_index * 3))
    return ReportSpec(title=f"{scenario} 장래년도 도시·교외간선도로 분석결과", groups=groups, rows=rows, widths=widths, row_categories=[sorted([r for r in future if _key(r)==key], key=lambda r:r.segment.year)[0].segment.road_category for key in base_keys], highlight_cells=highlights,header_spans=SEGMENT_HEADER_SPANS,row_pair_ids=[key[0] for key in base_keys],body_merge_columns=SEGMENT_MERGE_COLUMNS)


def comparison_report(
    project: ArterialProject,
    results: Iterable[AnalysisResult],
    year: int,
    before_scenario: str,
    after_scenario: str,
) -> ReportSpec:
    before = {_key(r): r for r in results if r.segment.scenario == before_scenario and r.segment.year == year}
    after = {_key(r): r for r in results if r.segment.scenario == after_scenario and r.segment.year == year}
    rows: list[list[str]] = []
    details: list[str] = []
    warnings: list[str] = []
    highlights: set[tuple[int, int]] = set()
    categories: list[str] = []
    ordered_keys = sorted(
        set(before) | set(after),
        key=lambda key: result_sort_key(after.get(key) or before[key]),
    )
    for key in ordered_keys:
        a, b = before.get(key), after.get(key)
        source = b or a
        if source is None:
            continue
        a_values = ["-", "-", "-"] if a is None else [
            _fmt_number(_report_volume(a)), _fmt_number(a.speed_kmh, 1), a.los
        ]
        b_values = ["-", "-", "-"] if b is None else [
            _fmt_number(_report_volume(b)), _fmt_number(b.speed_kmh, 1), b.los
        ]
        if a is None or b is None:
            change = ["-", "-", "-"]
            detail = "비교 대상 구간이 한쪽 상황에만 있습니다."
            warning = "비교구간 확인"
        else:
            a_volume, b_volume = _report_volume(a), _report_volume(b)
            volume_delta = b_volume - a_volume
            speed_delta = b.speed_kmh - a.speed_kmh
            volume_pct = "-" if a_volume == 0 else f"{volume_delta / a_volume * 100:+.1f}%"
            change = [
                _delta_text(volume_delta),
                _delta_text(speed_delta, 1),
                "-" if a.los == b.los else f"{a.los}→{b.los}",
            ]
            detail = f"교통량 증감률 {volume_pct} · V/c {a.vc_ratio:.2f}→{b.vc_ratio:.2f} · 지체 변화 {b.control_delay_s-a.control_delay_s:+.1f}초 · 차로수 {a.segment.lanes}→{b.segment.lanes} · g/C {a.segment.green_s/a.segment.cycle_s:.2f}→{b.segment.green_s/b.segment.cycle_s:.2f}"
            improved = b.speed_kmh > a.speed_kmh + 0.05 or ({"A":0,"B":1,"C":2,"D":3,"E":4,"F":5,"FF":6,"FFF":7}.get(b.los,9) < {"A":0,"B":1,"C":2,"D":3,"E":4,"F":5,"FF":6,"FFF":7}.get(a.los,9))
            warning = "검토 필요: 교통량 증가에도 평균통행속도/서비스수준 개선" if volume_delta > 0 and improved else ""
            if warning and (a.segment.manual_speed_kmh is not None or b.segment.manual_speed_kmh is not None): warning += " · 수동조정값 포함"
        row_index = len(rows)
        rows.append([*_segment_cells(source), *a_values, *b_values, *change, "상세"])
        categories.append(source.segment.road_category)
        details.append(detail)
        warnings.append(warning)
        if a and a.segment.manual_speed_kmh is not None: highlights.add((row_index, 7))
        if b and b.segment.manual_speed_kmh is not None: highlights.add((row_index, 10))
    short_a = "사업미시행시(A)" if before_scenario == "사업 미시행시" else f"{before_scenario}(A)"
    short_b = "사업시행시(B)" if after_scenario == "사업 시행시" else f"{after_scenario}(B)"
    groups = [
        SEGMENT_GROUP,
        (short_a, ["교통량\n(대/시)", "평균통행속도\n(km/h)", "서비스수준\n(LOS)"]),
        (short_b, ["교통량\n(대/시)", "평균통행속도\n(km/h)", "서비스수준\n(LOS)"]),
        ("변화(B-A)", ["교통량\n(대/시)", "평균통행속도\n(km/h)", "서비스수준\n(LOS)", "검토"]),
    ]
    return ReportSpec(
        title=f"{year}년 {before_scenario}·{after_scenario} 비교",
        groups=groups,
        rows=rows,
        widths=[*SEGMENT_WIDTHS,*([RESULT_WIDTH]*10)],
        row_categories=categories,
        highlight_cells=highlights,
        row_details=details,
        row_warnings=warnings,
        header_spans=SEGMENT_HEADER_SPANS,row_pair_ids=[key[0] for key in ordered_keys],body_merge_columns=SEGMENT_MERGE_COLUMNS,
    )


def report_for(project: ArterialProject, results: list[AnalysisResult], kind: str, year: int | None = None) -> ReportSpec:
    if kind == "현황 표":
        return current_report(project, results)
    if kind.startswith("장래년도 통합"):
        scenario = kind.split(" - ", 1)[1]
        return future_report(project, results, scenario)
    chosen_year = year or (project.future_years[0] if project.future_years else project.current_year)
    if kind == "미시행·시행 비교":
        return comparison_report(project, results, chosen_year, "사업 미시행시", "사업 시행시")
    return comparison_report(project, results, chosen_year, "사업 시행시", "개선대책 이행시")


def render_report_image(spec: ReportSpec, scale: float = 1.5, include_title: bool = True, compact: bool = False) -> QImage:
    margin, title_h, group_h, header_h, row_h = ((10, 0, 30, 44, 30) if compact else (24, 56 if include_title else 0, 42, 54, 38))
    width = margin * 2 + sum(spec.widths)
    height = margin * 2 + title_h + group_h + header_h + max(1, len(spec.rows)) * row_h
    image = QImage(int(width * scale), int(height * scale), QImage.Format_ARGB32)
    image.fill(QColor("white"))
    painter = QPainter(image)
    painter.scale(scale, scale)
    painter.setRenderHint(QPainter.Antialiasing)
    line = QPen(QColor("#9AA7B7"), 1)
    strong = QPen(QColor("#536273"), 1.4)
    font = QFont("Noto Sans KR", 10)
    painter.setFont(font)

    painter.setPen(QColor("#13233A"))
    if include_title:
        title_font = QFont(font); title_font.setPointSize(14); title_font.setBold(True); painter.setFont(title_font)
        painter.drawText(QRect(margin, margin, width - margin * 2, title_h), Qt.AlignCenter, spec.title)

    x, y = margin, margin + title_h
    painter.setPen(strong)
    header_offset = 0
    for group, headers in spec.groups:
        group_width = sum(spec.widths[header_offset:header_offset + len(headers)])
        painter.fillRect(QRect(x, y, group_width, group_h), QColor("#DCE4EC"))
        painter.drawRect(QRect(x, y, group_width, group_h))
        painter.setFont(QFont("Noto Sans KR", 9 if compact else 10, QFont.Bold))
        painter.drawText(QRect(x + 3, y + 2, group_width - 6, group_h - 4), Qt.AlignCenter | Qt.TextWordWrap, group)
        x += group_width
        header_offset += len(headers)

    y += group_h
    painter.setFont(QFont("Noto Sans KR", 8 if compact else 9, QFont.Bold))
    header_span_at={start:(span,label) for start,span,label in spec.header_spans}; covered=set()
    for column,(header,cell_width) in enumerate(zip(spec.headers,spec.widths)):
        if column in covered: continue
        span,label=header_span_at.get(column,(1,header)); width=sum(spec.widths[column:column+span]); covered.update(range(column+1,column+span))
        x=margin+sum(spec.widths[:column]); painter.fillRect(QRect(x,y,width,header_h),QColor("#EEF2F6")); painter.drawRect(QRect(x,y,width,header_h)); painter.drawText(QRect(x+4,y+3,width-8,header_h-6),Qt.AlignCenter|Qt.TextWordWrap,label)

    y += header_h
    painter.setFont(QFont("Noto Sans KR", 8 if compact else 9))
    body_rows = spec.rows or [["분석자료 없음", *([""] * (len(spec.widths) - 1))]]
    pair_spans: dict[int,int]={}
    if spec.row_pair_ids:
        index=0
        while index<len(spec.row_pair_ids):
            end=index+1
            while end<len(spec.row_pair_ids) and spec.row_pair_ids[end]==spec.row_pair_ids[index]: end+=1
            pair_spans[index]=end-index; index=end
    covered_body:set[tuple[int,int]]=set()
    for row_index, row in enumerate(body_rows):
        fill = QColor("#FFFFFF" if row_index % 2 == 0 else "#F8FAFC")
        for column,(value,cell_width) in enumerate(zip(row,spec.widths)):
            if (row_index,column) in covered_body: continue
            row_span=pair_spans.get(row_index,1) if column in spec.body_merge_columns else 1
            if row_span>1:
                for rr in range(row_index+1,row_index+row_span): covered_body.add((rr,column))
            x=margin+sum(spec.widths[:column]); height=row_h*row_span
            painter.fillRect(QRect(x, y, cell_width, height), fill)
            painter.setPen(line)
            painter.drawRect(QRect(x, y, cell_width, height))
            painter.setPen(QColor("#172033"))
            painter.drawText(QRect(x + 4, y + 2, cell_width - 8, height - 4), Qt.AlignCenter | Qt.TextWordWrap, str(value))
        y += row_h
    painter.end()
    image.setDevicePixelRatio(1.0)
    return image


def export_xlsx(spec: ReportSpec, path: str | Path) -> None:
    """앱 배포 환경에서 보고서 표를 엑셀 파일로 내보낸다."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "분석결과"
    sheet.sheet_view.showGridLines = False
    last_col = len(spec.widths)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    sheet.cell(1, 1, spec.title)
    sheet.cell(1, 1).font = Font(name="Noto Sans KR", size=14, bold=True)
    sheet.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 30
    thin = Side(style="thin", color="9AA7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    col = 1
    for group, headers in spec.groups:
        start = col
        end = col + len(headers) - 1
        if end > start:
            sheet.merge_cells(start_row=2, start_column=start, end_row=2, end_column=end)
        sheet.cell(2, start, group)
        for index, header in enumerate(headers, start=start):
            sheet.cell(3, index, header)
        col = end + 1
    for start,span,label in spec.header_spans:
        excel_start=start+1; excel_end=excel_start+span-1
        if excel_end>excel_start: sheet.merge_cells(start_row=3,start_column=excel_start,end_row=3,end_column=excel_end)
        sheet.cell(3,excel_start,label)
    def excel_value(value: str):
        text = str(value).strip()
        if re.fullmatch(r"[+-]?\d[\d,]*(?:\.\d+)?", text):
            number = float(text.replace(",", ""))
            return int(number) if number.is_integer() else number
        return text

    for row_index, row in enumerate(spec.rows, start=4):
        for col_index, value in enumerate(row, start=1):
            sheet.cell(row_index, col_index, excel_value(value))
    for row in sheet.iter_rows(min_row=2, max_row=max(3, 3 + len(spec.rows)), min_col=1, max_col=last_col):
        for cell in row:
            cell.font = Font(name="Noto Sans KR", size=9, bold=cell.row <= 3)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
            if cell.row == 2:
                cell.fill = PatternFill("solid", fgColor="DCE4EC")
            elif cell.row == 3:
                cell.fill = PatternFill("solid", fgColor="EEF2F6")
    for index, pixel_width in enumerate(spec.widths, start=1):
        sheet.column_dimensions[chr(64 + index) if index <= 26 else f"A{chr(64 + index - 26)}"].width = max(10, pixel_width / 8)
    sheet.freeze_panes = "C4"
    workbook.save(path)
