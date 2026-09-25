"""기존 가로분석 Excel과 교차로 정리표를 .ara1 프로젝트로 변환한다."""

from __future__ import annotations

import argparse
import re
from datetime import datetime
from pathlib import Path
from uuid import NAMESPACE_URL, uuid5

from openpyxl import load_workbook

from arterial_analysis.engine import ROAD_CATEGORIES, SegmentInput
from arterial_analysis.excel_links import normalize_expression, read_saved_expressions
from arterial_analysis.intersection_import import match_intersection_los, read_intersection_los
from arterial_analysis.project import ArterialProject


SCENARIO_BY_PREFIX = {
    "현황": "현황",
    "미시행": "사업 미시행시",
    "시행": "사업 시행시",
}


def parse_sheet(sheet_name: str) -> tuple[str, int]:
    year = re.search(r"20\d{2}", sheet_name)
    scenario = next((value for prefix, value in SCENARIO_BY_PREFIX.items() if sheet_name.startswith(prefix)), None)
    if not year or not scenario:
        raise ValueError(f"분석 시트명을 해석할 수 없습니다: {sheet_name}")
    return scenario, int(year.group())


def split_intersection(value: object) -> tuple[str, str]:
    text = str(value or "").strip()
    match = re.match(r"^([^.．]+)[.．]\s*(.+)$", text)
    if not match:
        return "", text
    number, name = match.groups()
    number = str(int(number)) if number.isdigit() else number.upper()
    return number, name.strip()


def formula_reference(formula: object) -> tuple[str, str]:
    text = str(formula or "")
    matches = re.findall(r"\[[^\]]+\]([^'!]+)'?!\$?([A-Za-z]{1,3})\$?([1-9]\d*)", text)
    if not matches:
        raise ValueError(f"외부 셀 참조식을 찾을 수 없습니다: {text}")
    sheets = {sheet for sheet, _, _ in matches}
    if len(sheets) != 1:
        raise ValueError(f"한 연결식에 여러 시트가 포함되어 있습니다: {text}")
    return matches[0][0], normalize_expression(",".join(f"{column}{row}" for _, column, row in matches))


def merged_value(worksheet, row: int, column: int):
    value = worksheet.cell(row, column).value
    if value not in (None, ""):
        return value
    coordinate = worksheet.cell(row, column).coordinate
    for merged in worksheet.merged_cells.ranges:
        if coordinate in merged:
            return worksheet.cell(merged.min_row, merged.min_col).value
    return value


def numeric(value: object, label: str) -> float:
    if value in (None, ""):
        raise ValueError(f"{label} 값이 비어 있습니다.")
    return float(value)


def linked_value(path: Path, sheet: str, expression: str, kind: str) -> float:
    value, _, error = read_saved_expressions(str(path), sheet, [expression], kind)[expression]
    if error:
        raise ValueError(f"{path.name} · {sheet}!{expression}: {error}")
    return value


def build_project(main_path: Path, traffic_path: Path, intersection_path: Path) -> ArterialProject:
    formulas = load_workbook(main_path, data_only=False, keep_links=False)
    values = load_workbook(main_path, data_only=True, keep_links=False)
    try:
        source_sheets = [name for name in formulas.sheetnames if re.match(r"^(현황|미시행|시행)20\d{2}$", name)]
        if not source_sheets:
            raise ValueError("현황·미시행·시행 분석 시트를 찾을 수 없습니다.")
        parsed = [(name, *parse_sheet(name)) for name in source_sheets]
        current_year = next(year for _, scenario, year in parsed if scenario == "현황")
        future_years = sorted({year for _, scenario, year in parsed if scenario != "현황"})
        project = ArterialProject(name="대구미래 가로분석", current_year=current_year, future_years=future_years)
        project.analysis_tabs = []
        intersection_groups = {(group.scenario, group.year): group for group in read_intersection_los(intersection_path)}
        for sheet_name, scenario, year in parsed:
            ws_formula, ws_value = formulas[sheet_name], values[sheet_name]
            tab_volume_sheet = None
            result_offset = 16 if scenario == "사업 시행시" else 14
            rows = [row for row in range(5, 17) if ws_formula.cell(row, 3).value in ("→", "←")]
            for row in rows:
                road_name = str(merged_value(ws_formula, row, 1) or "").strip()
                start_number, start_name = split_intersection(merged_value(ws_formula, row, 2))
                end_number, end_name = split_intersection(merged_value(ws_formula, row, 4))
                direction = str(ws_formula.cell(row, 3).value)
                volume_sheet, volume_expression = formula_reference(ws_formula.cell(row, 14).value)
                tab_volume_sheet = tab_volume_sheet or volume_sheet
                volume = linked_value(traffic_path, volume_sheet, volume_expression, "volume")
                phf_formula = ws_formula.cell(row, 15).value
                if isinstance(phf_formula, str) and phf_formula.startswith("="):
                    phf_sheet, phf_expression = formula_reference(phf_formula)
                    phf_path = traffic_path
                else:
                    phf_sheet, phf_expression = sheet_name, f"O{row}"
                    phf_path = main_path
                phf = linked_value(phf_path, phf_sheet, phf_expression, "phf")
                arrival_number, arrival_name = (end_number, end_name) if direction == "→" else (start_number, start_name)
                group = intersection_groups.get((scenario, year))
                record = match_intersection_los(group.records, arrival_number, arrival_name) if group else None
                speed_text = str(ws_value.cell(row + result_offset, 10).value or "")
                speed_match = re.search(r"\d+(?:\.\d+)?", speed_text)
                physical_key = "|".join((road_name, start_number or start_name, end_number or end_name))
                uid_key = "|".join((scenario, str(year), physical_key, direction))
                segment = SegmentInput(
                    uid=uuid5(NAMESPACE_URL, uid_key).hex,
                    comparison_id=uuid5(NAMESPACE_URL, physical_key).hex,
                    scenario=scenario,
                    year=year,
                    road_category=ROAD_CATEGORIES[0],
                    road_name=road_name,
                    start_number=start_number,
                    start_name=start_name,
                    direction=direction,
                    end_number=end_number,
                    end_name=end_name,
                    length_km=numeric(ws_value.cell(row, 5).value, f"{sheet_name}!E{row}") / 1000,
                    lanes=int(round(numeric(ws_value.cell(row, 17).value, f"{sheet_name}!Q{row}"))),
                    functional_class="저규격",
                    road_condition_override="보통",
                    arterial_type_override="자동",
                    cycle_s=numeric(ws_value.cell(row, 11).value, f"{sheet_name}!K{row}"),
                    green_s=numeric(ws_value.cell(row, 12).value, f"{sheet_name}!L{row}"),
                    main_volume=volume,
                    phf=phf,
                    intersection_los=record.los if record else "",
                    speed_limit_kmh=float(speed_match.group()) if speed_match else None,
                    blank_fields=[],
                    volume_source_cell=volume_expression,
                    volume_source_path=str(traffic_path.resolve()),
                    volume_source_sheet=volume_sheet,
                    volume_link_status="ok",
                    volume_last_value=volume,
                    phf_source_cell=phf_expression,
                    phf_source_path=str(phf_path.resolve()),
                    phf_source_sheet=phf_sheet,
                    phf_link_status="ok",
                    phf_last_value=phf,
                )
                project.segments.append(segment)
            group = intersection_groups.get((scenario, year))
            project.analysis_tabs.append({
                "scenario": scenario,
                "year": year,
                "source_path": str(traffic_path.resolve()),
                "source_sheet": tab_volume_sheet or "",
                "memo": "기존 가로분석 Excel에서 원본 셀 연결과 함께 변환",
                "intersection_los_source_path": str(intersection_path.resolve()),
                "intersection_los_source_sheet": group.sheet if group else "",
                "intersection_los_source_header": group.header if group else "",
            })
        project.source_note = f"가로분석: {main_path.name} · 교통량/PHF 원본: {traffic_path.name} · 교차로 LOS: {intersection_path.name}"
        project.add_log({
            "at": datetime.now().astimezone().isoformat(timespec="seconds"),
            "kind": "change",
            "scenario": "",
            "year": 0,
            "segment": "",
            "uid": "",
            "field": "Excel 프로젝트 변환",
            "old": "",
            "new": f"{len(project.segments)}개 방향 구간",
            "message": "교통량·PHF 원본 셀, 교차로 LOS 및 제한속도 참고값 포함",
        })
        return project
    finally:
        formulas.close()
        values.close()


def validate_project(project: ArterialProject) -> None:
    expected={("현황",2026):10,("사업 미시행시",2035):10,("사업 시행시",2035):12}
    actual={key:len(project.rows(*key)) for key in project.tab_keys()}
    if actual != expected:
        raise ValueError(f"변환 행 수 불일치: {actual}")
    incomplete=[segment.uid for segment in project.segments if not project.is_complete(segment)]
    if incomplete:
        raise ValueError(f"입력 미완료 구간: {len(incomplete)}개")
    missing_speed=[segment.uid for segment in project.segments if segment.speed_limit_kmh is None]
    missing_los=[segment.uid for segment in project.segments if not segment.intersection_los]
    if missing_speed or missing_los:
        raise ValueError(f"참고값 누락: 제한속도 {len(missing_speed)}개, 교차로 LOS {len(missing_los)}개")
    if any(segment.road_category == ROAD_CATEGORIES[2] for segment in project.segments):
        raise ValueError("대구 변환 프로젝트에는 내부도로가 포함되면 안 됩니다.")
    results=project.analyze_all()
    if len(results) != len(project.segments) or any(result.arterial_type != "유형 III" for result in results):
        raise ValueError("전체 구간이 유형 III으로 분석되지 않았습니다.")


def main() -> None:
    parser=argparse.ArgumentParser()
    parser.add_argument("main",type=Path)
    parser.add_argument("traffic",type=Path)
    parser.add_argument("intersection",type=Path)
    parser.add_argument("output",type=Path)
    args=parser.parse_args()
    project=build_project(args.main.resolve(),args.traffic.resolve(),args.intersection.resolve())
    validate_project(project);args.output.parent.mkdir(parents=True,exist_ok=True);project.save(args.output)
    loaded=ArterialProject.load(args.output);validate_project(loaded)
    print(f"saved={args.output.resolve()}")
    print(f"tabs={loaded.tab_keys()}")
    print(f"segments={len(loaded.segments)} results={len(loaded.analyze_all())}")
    print(f"volume_links={sum(bool(row.volume_source_cell) for row in loaded.segments)} phf_links={sum(bool(row.phf_source_cell) for row in loaded.segments)}")
    print(f"intersection_los={sum(bool(row.intersection_los) for row in loaded.segments)} speed_limits={sum(row.speed_limit_kmh is not None for row in loaded.segments)}")


if __name__ == "__main__":
    main()
